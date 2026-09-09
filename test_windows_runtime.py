"""Portable acceptance checks with disposable data; never modify the local workspace DB."""
import base64
import io
import json
from pathlib import Path
import sqlite3
import tempfile
import threading
import urllib.request
from http.server import ThreadingHTTPServer

import fitz
from openpyxl import load_workbook
import server
from start_local import prepare_data


def run():
    with tempfile.TemporaryDirectory(prefix='kingpolymer-windows-') as temp:
        server.DATA = Path(temp) / '中文路径' / 'data'
        server.BACKUPS = Path(temp) / 'backups'
        server.DB_PATH = server.DATA / 'guangjun_screw.db'
        server.AUTH_USER = server.AUTH_PASSWORD = ''
        server.AUTH_REQUIRED = False
        prepare_data(server.DATA)
        assert not (server.DATA / 'web_auth.db').exists()
        server.init_db()
        class QuietHandler(server.Handler):
            def log_message(self, *args):
                pass
        http = ThreadingHTTPServer(('127.0.0.1', 0), QuietHandler)
        thread = threading.Thread(target=http.serve_forever, daemon=True)
        thread.start()
        url = f'http://127.0.0.1:{http.server_port}'
        opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))

        def request(path, body=None):
            req = urllib.request.Request(url + path,
                data=None if body is None else json.dumps(body).encode(),
                headers={'Content-Type': 'application/json'})
            with opener.open(req, timeout=10) as response:
                content = response.read()
                return json.loads(content) if 'application/json' in response.headers['Content-Type'] else content

        try:
            assert b'v1.28' in request('/')
            data = request('/api/bootstrap')
            assert len(data['components']) >= 70 and len(data['templates']) >= 20
            assert not data['auth']['enabled']
            template = next(t for t in data['templates'] if t['machine'] == '50' and t['is_default'])
            project = {'machine': '50', 'sequence': template['sequence'], 'ports': {},
                       'metadata': {'drawing_name': 'Windows验收'},
                       'override_reason': 'Isolated acceptance test', 'confirm_code': server.CONFIRM_CODE}
            saved = request('/api/projects', project)
            assert saved['ok']
            assert request('/api/projects/' + saved['id'])['metadata']['drawing_name'] == 'Windows验收'
            content = request('/api/export/xlsx', project)
            workbook = load_workbook(io.BytesIO(content))
            assert '安装顺序' in workbook.sheetnames and '机筒配置' in workbook.sheetnames
            workbook.close()
            imported = request('/api/import', {'filename': 'Windows验收.xlsx', 'data': base64.b64encode(content).decode()})
            assert set(imported['sequence']) == set(project['sequence'])
            with fitz.open() as pdf:
                page = pdf.new_page()
                page.insert_text((40, 40), 'GFA-2-60-60')
                content = pdf.tobytes()
            imported = request('/api/import', {'filename': 'test.pdf', 'data': base64.b64encode(content).decode()})
            assert imported['sequence'] == ['GFA-2-60-60']
            backup = request('/api/backup', {})
            assert Path(backup['file']).is_file()
            with server.db() as connection:
                assert connection.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
            prepare_data(server.DATA)
            server.init_db()
            assert request('/api/projects/' + saved['id'])['metadata']['drawing_name'] == 'Windows验收'
        finally:
            http.shutdown()
            http.server_close()
            thread.join()
    print('PASS: Chinese paths, isolated data copy, project save/reload, Excel export/import, PDF import, backups, repeated initialization and Windows file cleanup')


if __name__ == '__main__':
    run()
