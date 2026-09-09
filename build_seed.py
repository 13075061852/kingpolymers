#!/usr/bin/env python3
"""Build the initial component/template library from Guangjun's archived drawings.
This script is for maintenance; the generated data/seed.json ships with the app.
It never modifies source files.
"""
from __future__ import annotations
import json, os, re, glob, hashlib
from datetime import datetime
from pathlib import Path

ROOT = "/Users/lsy/Desktop/广俊塑料AI/13.螺杆组合"
ONEDRIVE_DRAWINGS = "/Users/lsy/Library/CloudStorage/OneDrive-个人/文档/挤出机/莱斯特瑞兹/螺杆组合"
SOURCE_DIRS = {
    "50": [os.path.join(ROOT, "50CC PRO"), os.path.join(ONEDRIVE_DRAWINGS, "ZSE50螺杆组合")],
    "60": [os.path.join(ROOT, "60CC PRO"), os.path.join(ONEDRIVE_DRAWINGS, "ZES60螺杆组合")],
}
OUT = os.path.join(os.path.dirname(__file__), "data", "seed.json")


def element_parameters(name: str):
    """Decode Leistritz model numbers without changing the source model name."""
    numeric_name = re.sub(r'\(\d+\)$', '', name or '')
    nums = [int(x) for x in re.findall(r"\d+", numeric_name)]
    kind = element_type(name)
    length = nums[-1] if nums else 0
    pitch = discs = angle = None
    if kind in ("KB", "KBX") and len(nums) >= 4:
        discs, length, angle = nums[0], nums[-2], nums[-1]
    elif kind in ("GFA", "GFF", "GFM", "SME") and len(nums) >= 3:
        pitch = nums[-2]
    direction = "LI" if re.search(r"(?:-|°|度|^)LI?$", name, re.I) else ("RE" if re.search(r"(?:-|°|度|^)RE$", name, re.I) else "")
    return {"length": length, "pitch": pitch, "discs": discs, "angle": angle, "direction": direction}


def element_length(name: str) -> int:
    return element_parameters(name)["length"]


def element_type(name: str) -> str:
    up = name.upper()
    for kind in ("GFA", "GFF", "GFM", "SME", "KBX", "KB", "KS", "SPACER"):
        if up.startswith(kind):
            return kind.title() if kind == "SPACER" else kind
    return "OTHER"


def pdf_sequence(path: str):
    import fitz
    doc = fitz.open(path)
    if len(doc) < 3:
        return []
    words = doc[2].get_text("words")
    candidates = []
    pattern = re.compile(r"^(?:GFA|GFF|GFM|SME|KBX?|KS|Spacer)[A-Za-z0-9°()/_+.\-]*$", re.I)
    for w in words:
        x, y, _, _, text = w[:5]
        if y < 340 and x < 410 and pattern.match(text) and "configuration" not in text.lower():
            candidates.append((x, y, text))
    if not candidates:
        # Fallback for structurally different exports.
        lines = doc[2].get_text("text").splitlines()
        return [s.strip() for s in lines if pattern.match(s.strip())]
    # SIGMA tables use multiple vertical columns. Sort columns from left to right,
    # and each column from top to bottom.
    xs = sorted(x for x, _, _ in candidates)
    groups = []
    for item in sorted(candidates, key=lambda z: z[0]):
        for group in groups:
            if abs(group[0][0] - item[0]) < 45:
                group.append(item)
                break
        else:
            groups.append([item])
    groups.sort(key=lambda g: min(x for x, _, _ in g))
    result = []
    for group in groups:
        result.extend(text for _, _, text in sorted(group, key=lambda z: z[1]))
    return result


def pdf_catalog(path: str):
    import fitz
    doc = fitz.open(path)
    if len(doc) < 2:
        return []
    pattern = re.compile(r"^(?:GFA|GFF|GFM|SME|KBX?|KS|Spacer)[A-Za-z0-9°()/_+.\-]*$", re.I)
    return [line.strip() for line in doc[1].get_text("text").splitlines() if pattern.match(line.strip())]


def pdf_date(path: str):
    import fitz
    raw = fitz.open(path).metadata.get("creationDate", "")
    m = re.search(r"D:(\d{4})(\d{2})(\d{2})", raw)
    return f"{m.group(1)}-{m.group(2)}-{m.group(3)}" if m else ""


def version_from_name(name: str):
    m = re.search(r"48D_+([^_]+)_", name, re.I)
    return m.group(1) if m else os.path.splitext(name)[0]


def material_from_name(name: str):
    stem = os.path.splitext(name)[0]
    m = re.search(r"Guang Jun_+(.+)$", stem, re.I)
    return m.group(1).replace("(1)", "") if m else ""


def component_record(machine, name):
    params = element_parameters(name)
    return {
        "id": hashlib.sha1(f"{machine}:{name}".encode()).hexdigest()[:16],
        "machine": machine,
        "name": name,
        "type": element_type(name),
        **params,
        "stock": None,
        "note": "由莱斯特瑞兹历史图纸导入"
    }


def main():
    templates, catalog = [], {}
    sequence_seen = set()
    for machine, folders in SOURCE_DIRS.items():
        paths = sorted({p for folder in folders for p in glob.glob(os.path.join(folder, "*.pdf"))})
        for path in paths:
            if os.path.basename(path).startswith("Schnecken"):
                continue
            try:
                seq = pdf_sequence(path)
                for n in pdf_catalog(path) + seq:
                    if element_length(n) > 0:
                        catalog[(machine, n)] = component_record(machine, n)
                if not seq:
                    continue
                filename = os.path.basename(path)
                date = pdf_date(path)
                version = version_from_name(filename)
                signature = (machine, version, date, tuple(seq))
                if signature in sequence_seen:
                    continue
                sequence_seen.add(signature)
                if machine == "50" and version.upper() == "1G":
                    date = "2025-03-25"
                templates.append({
                    "id": hashlib.sha1((machine + filename + date + "|".join(seq)).encode()).hexdigest()[:16],
                    "machine": machine,
                    "name": f"{version} · {date} · {material_from_name(filename)}",
                    "version": version,
                    "date": date,
                    "material": material_from_name(filename),
                    "source": filename,
                    "sequence": seq,
                    "total_length": sum(element_length(n) for n in seq),
                    "is_default": version.upper() == "1G"
                })
            except Exception as exc:
                print("skip", path, exc)

    # Include the operator-maintained Excel histories as additional 50CC templates.
    try:
        from openpyxl import load_workbook
        xlsx_seen = set()
        for book in (os.path.join(ROOT, "B线螺杆组合.xlsx"),
                     "/Users/lsy/Library/CloudStorage/OneDrive-个人/万洋工厂/13.螺杆组合/广俊螺杆组合.xlsx"):
            if not os.path.exists(book):
                continue
            wb = load_workbook(book, read_only=True, data_only=True)
            for ws in wb.worksheets:
                seq = []
                note = ""
                for row in ws.iter_rows(values_only=True):
                    val = str(row[0]).strip() if row and row[0] is not None else ""
                    if re.match(r"^(?:GFA|GFF|GFM|SME|KBX?|KS|Spacer)", val, re.I):
                        seq.append(val)
                        catalog[("50", val)] = component_record("50", val)
                    elif val and not note:
                        note = val
                signature = ("50", tuple(seq))
                if seq and signature not in xlsx_seen:
                    xlsx_seen.add(signature)
                    templates.append({
                        "id": hashlib.sha1((book + ws.title).encode()).hexdigest()[:16],
                        "machine": "50", "name": f"Excel · {ws.title} · {Path(book).stem}", "version": ws.title,
                        "date": "", "material": "", "source": os.path.basename(book),
                        "sequence": seq, "total_length": sum(element_length(n) for n in seq),
                        "is_default": False, "note": note
                    })
    except Exception as exc:
        print("xlsx histories skipped", exc)

    data = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "machines": {
            "50": {
                "id": "50", "name": "ZSE 50CC Pro", "machine_no": "04037198",
                "diameter": 50, "center_distance": 40, "rotation": "right",
                "section_length": 210, "sections": 12, "barrel_length": 2548,
                "element_length": 2590, "entry_offset": 42, "position_origin": 0, "flange_length": 27,
                "heat_barrier": 1, "side_model": "LSB047",
                "barrel_configuration": [
                    {"pos":"A","mm":210,"name":"Zylinder E","role":"feed"},
                    {"pos":"B","mm":211,"name":"Wärmesperre","role":"heat"},
                    {"pos":"C","mm":421,"name":"Zylinder 0","role":"barrel"},
                    {"pos":"D","mm":631,"name":"Zylinder 0","role":"barrel"},
                    {"pos":"E","mm":841,"name":"Zylinder 1","role":"natural","port_key":"natural4"},
                    {"pos":"F","mm":1051,"name":"Zylinder S (LSB047)","role":"side"},
                    {"pos":"G","mm":1261,"name":"Zylinder 0","role":"barrel"},
                    {"pos":"H","mm":1471,"name":"Zylinder 1","role":"natural","port_key":"natural7"},
                    {"pos":"I","mm":1681,"name":"Zylinder S (LSB047)","role":"side"},
                    {"pos":"J","mm":1891,"name":"Zylinder 0","role":"barrel"},
                    {"pos":"K","mm":2101,"name":"Zylinder 0","role":"barrel"},
                    {"pos":"L","mm":2311,"name":"Zylinder 1","role":"vacuum"},
                    {"pos":"M","mm":2521,"name":"Zylinder 0","role":"barrel"},
                    {"pos":"N","mm":2548,"name":"Zwischenflansch","role":"flange"}
                ]
            },
            "60": {
                "id": "60", "name": "ZSE 60CC", "machine_no": "05047144",
                "diameter": 60, "center_distance": 50, "rotation": "left",
                "section_length": 240, "sections": 12, "barrel_length": 2910,
                "element_length": 2970, "entry_offset": 60, "position_origin": 1, "flange_length": 30,
                "heat_barrier": 0, "side_model": "LSB56",
                "barrel_configuration": [
                    {"pos":"A","mm":240,"name":"Zyl-E","role":"feed"},
                    {"pos":"B","mm":480,"name":"Zyl-0","role":"barrel"},
                    {"pos":"C","mm":720,"name":"Zyl-0","role":"barrel"},
                    {"pos":"D","mm":960,"name":"Zyl-1","role":"natural","port_key":"natural4"},
                    {"pos":"E","mm":1200,"name":"Zyl-S (LSB 56)","role":"side"},
                    {"pos":"F","mm":1440,"name":"Zyl-0","role":"barrel"},
                    {"pos":"G","mm":1680,"name":"Zyl-1","role":"natural","port_key":"natural7"},
                    {"pos":"H","mm":1920,"name":"Zyl-S (LSB 56)","role":"side"},
                    {"pos":"I","mm":2160,"name":"Zyl-0","role":"barrel"},
                    {"pos":"J","mm":2400,"name":"Zyl-0","role":"barrel"},
                    {"pos":"K","mm":2640,"name":"Zyl-1","role":"vacuum"},
                    {"pos":"L","mm":2880,"name":"Zyl-0","role":"barrel"},
                    {"pos":"M","mm":2910,"name":"Zwischenflansch","role":"flange"}
                ]
            }
        },
        "components": sorted(catalog.values(), key=lambda c: (c["machine"], c["type"], c["length"], c["name"])),
        "templates": sorted(templates, key=lambda t: (t["machine"], t["date"], t["version"]))
    }
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
    print(f"wrote {OUT}: {len(data['components'])} components, {len(templates)} templates")

if __name__ == "__main__":
    main()
