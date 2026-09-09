"""Barrel layout parity, safety, persistence and export tests (temporary DB only)."""
import io,json,subprocess,tempfile,threading,urllib.request,urllib.error
from pathlib import Path
from http.server import ThreadingHTTPServer
from openpyxl import load_workbook
import server,barrel_models

def run():
 with tempfile.TemporaryDirectory() as directory:
  server.DATA=Path(directory)/'data';server.BACKUPS=Path(directory)/'backups';server.DB_PATH=server.DATA/'test.db';server.init_db()
  specs=server.machine_specs()
  for machine,spec in specs.items():
   assert all(x['machine']==machine and x['nominal_diameter']==int(machine) and x['bore_diameter'] is None for x in spec['barrel_catalog'])
   assert not any(x['role']=='heat' for x in specs['60']['barrel_catalog'])
   modules=[dict(uid=f'm{i}',key=f'{machine}:{role}',open=True) for i,role in enumerate(['feed','barrel','natural','side','vacuum','flange'])]
   ports={'natural4':True,'natural7':True,'barrel_layout':dict(version=1,machine=machine,modules=modules)}
   rows=server.barrel_configuration(machine,ports);assert rows[1]['length']==spec['section_length']
   assert rows[-1]['mm']==5*spec['section_length']+spec['flange_length']
   zones=server.restriction_zones(machine,ports);assert len(zones)==3
   assert zones[0]['start']==2.5*spec['section_length']-({'50':72,'60':80}[machine]/2)-30
   ports['barrel_layout']['modules'].insert(0,dict(uid='extra',key=f'{machine}:barrel',open=True))
   assert server.restriction_zones(machine,ports)[0]['start']==zones[0]['start']+spec['section_length']
   ports['barrel_layout']['modules'][3]['open']=False
   assert len(server.restriction_zones(machine,ports))==2
   # Node runs the real frontend model on the same server-provided catalog.
   program="const M=(await import('./frontend/src/domain/barrel-models.js')).default;const x=JSON.parse((await import('node:fs')).readFileSync(0,'utf8'));console.log(JSON.stringify({rows:M.rows(x.machine,x.spec,x.ports),zones:M.zones(x.machine,x.spec,x.ports),warnings:M.warnings(x.machine,x.spec,x.ports)}))"
   result=json.loads(subprocess.check_output(['node','--input-type=module','-e',program],input=json.dumps(dict(machine=machine,spec=spec,ports=ports)).encode(),cwd=server.ROOT))
   assert result['zones']==server.restriction_zones(machine,ports)
   assert result['warnings']==barrel_models.warnings(machine,spec,ports)
   for js,py in zip(result['rows'],server.barrel_configuration(machine,ports)):
    for key in ('length','start','mm','role','machine','nominal_diameter','open','annotation'):assert js[key]==py[key],key
   for bad in [dict(version=1,machine='other',modules=[]),dict(version=1,machine=machine,modules=[dict(uid='x',key=('60' if machine=='50' else '50')+':barrel',open=True)]),dict(version=1,machine=machine,modules=[dict(uid='x',key=f'{machine}:barrel',open='false')])]:
    try:barrel_models.normalize(machine,spec,bad)
    except ValueError:pass
    else:raise AssertionError('invalid layout accepted')
  empty={'barrel_layout':dict(version=1,machine='60',modules=[])}
  assert server.barrel_configuration('60',empty)==[]
  assert server.restriction_zones('60',empty)==[]
  assert len(barrel_models.warnings('60',specs['60'],empty))>=4
  assert barrel_models.position(26)=='AA'
  httpd=ThreadingHTTPServer(('127.0.0.1',0),server.Handler);thread=threading.Thread(target=httpd.serve_forever,daemon=True);thread.start()
  def request(path,data=None):
   req=urllib.request.Request(f'http://127.0.0.1:{httpd.server_port}'+path,data=None if data is None else json.dumps(data).encode(),headers={'Content-Type':'application/json'})
   try:
    with urllib.request.urlopen(req) as r:return r.status,r.read()
   except urllib.error.HTTPError as e:return e.code,e.read()
  try:
   payload=dict(machine='60',sequence=['GFA-2-60-60'],ports=ports,metadata={'drawing_name':'temporary barrel test'})
   # Use a 60 layout, ensure JSON round trip and no silent stripping of modules.
   assert payload['ports']['barrel_layout']['machine']=='60'
   code,data=request('/api/projects',payload);assert code==409
   payload.update(override_reason='test engineering approval',confirm_code=server.CONFIRM_CODE)
   code,data=request('/api/projects',payload);assert code==200,data;pid=json.loads(data)['id']
   code,data=request('/api/projects/'+pid);loaded=json.loads(data);assert loaded['ports']==server.normalize_ports('60',ports)
   parsed=server.parse_import('project.json',json.dumps({'project':loaded}).encode());assert parsed['ports']==loaded['ports']
   code,data=request('/api/export/xlsx',{'project':loaded});assert code==200
   wb=load_workbook(io.BytesIO(data),data_only=True);assert wb['机筒配置'].max_row==len(ports['barrel_layout']['modules'])+1
   assert wb['机筒配置']['H2'].value==60
   assert wb['工程复核'].max_row>2
   code,data=request('/api/projects/'+pid+'/release',{'reason':'unapproved'});assert code==409
   cross=dict(payload,ports={'barrel_layout':dict(version=1,machine='50',modules=[])})
   code,data=request('/api/projects',cross);assert code==400
   # Existing projects with no barrel layout still use the original standard.
   assert len(server.barrel_configuration('50',{}))==14
   assert len(server.barrel_configuration('60',{}))==13
  finally:httpd.shutdown();httpd.server_close();thread.join()
 print('PASS: machine dimensions/isolation, JS/Python parity, moving safety zones, approvals, project/JSON/XLSX round trips, empty layouts, legacy defaults')
if __name__=='__main__':run()
